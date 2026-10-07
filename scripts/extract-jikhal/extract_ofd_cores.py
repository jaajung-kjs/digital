"""직할 OFD 선번장 → 국·판(슬롯)별 코어 용도·상태·비고 (`backend/src/seed/jikhalGangwonCores.ts`).

    python3 -I scripts/extract-jikhal/extract_ofd_cores.py 직할/직할_OFD선번장_260203.xlsx backend/src/seed/jikhalGangwonCores.ts

판 나누기는 `jikhalGangwonLines.ts` 머리말의 규칙 그대로다: 국 시트 안에서 「상대국명」이 새로 적힌
코어 행이 판 하나의 시작이다(48심 블록은 판 두 장, 24심 블록 안의 12심 판도 그렇다). 판 안의 코어
번호는 `No. − 판 첫 No. + 1`. 뽑은 판의 상대국·심수는 `jikhalGangwon.ts` 의 `PLATES` 와, 선로명은
`jikhalGangwonLines.ts` 의 `PLATE_LINES` 와 맞아야 한다 — 어긋나면 실패한다(조용히 넘기지 않는다).

코어 행(열: No.·상대국명·유형·회선수용내역·비고 … 결과):
- 상태 글: 유형·회선수용내역·비고 셀(과 결과 셀이 「단선」일 때)에 `단선` → broken, `사용불가` → unusable.
  `측정불가(…)` 도 용도가 아니라 상태 글이다(상태는 ok). 상태 글 원문은 note 앞에 둔다.
- purpose = 유형 · 회선수용내역 (상태 글인 셀과 빈 셀은 뺀다).
- note    = 상태 글 원문 · 비고 (비고가 그 판의 선로명 셀과 같으면 뺀다).
- 비고의 `(\\d+)~(\\d+)코어` + 사용불가 는 그 판의 코어 범위 전부를 unusable + 그 원문 note 로 싣는다.
- role    = 회선수용내역 셀에 **명시 표기**가 있을 때만 — 주회선 `㈜`·`(주)`·`주회선`·`-주)`·`주_`,
  예비회선 `(예)`·`예비회선`·`-예비)`·`예비_`·`예비(`. 둘 다 있으면 실패(추측하지 않는다). 표기가 없으면 null.
  `I-MUX(B)` 의 `(B)` 는 장치 이름(I-MUX(A)/(B))이라 표기가 아니다 — 대국 끝은 같은 코어를 「주_Tx」 로 적는다.
  `본사-주From본부계 OFD 43~46` 은 네 코어 모두 「주」 인데 그중 44·46 은 예비 코어(본사-예비·동해-예비)에서
  온다 — 코어별 표기가 아니라 묶음 이름이라 표기로 보지 않는다(정규식이 `-주)` 만 잡는다).
  코어 하나의 role 은 두 끝 공유 — 두 끝을 합치는 건 시드(`jikhal.ts`)가 한다.
셀 글은 공백(줄바꿈 포함)을 한 칸으로 줄인다. 값이 모두 비고 상태 ok 인 코어는 싣지 않는다.
"""
import re
import sys

import openpyxl

SHEET_KEYS = {
    '(구)춘천': 'guchuncheon', '(신)춘천': 'sinchuncheon', '북춘천': 'bukchuncheon', '남춘천': 'namchuncheon',
    '서홍천': 'seohongcheon', '열병합': 'yeolbyeonghap', '인제': 'inje', '양구': 'yanggu', '철원': 'cheorwon',
    '화천': 'hwacheon', '화천HP': 'hwacheonhp', '춘천HP': 'chuncheonhp',
}
COL_NO, COL_PEER, COL_KIND, COL_USE, COL_NOTE, COL_RESULT, COL_LINE = 0, 1, 2, 3, 4, 9, 10
HEADER = ['No.', '상대국명', '유형', '회선수용내역', '비고']
RANGE = re.compile(r'(\d+)~(\d+)코어')
ROLE_MAIN = re.compile(r'㈜|\(주\)|주회선|-주\)|주_')
ROLE_BACKUP = re.compile(r'\(예\)|예비회선|-예비\)|예비_|예비\(')


def fail(msg):
    sys.exit(f'[extract_ofd_cores] 실패: {msg}')


def clean(v):
    if v is None:
        return None
    s = ' '.join(str(v).split())
    return s or None


def line_name(v):
    """`jikhalGangwonLines.ts` 의 선로명 규칙 — 줄은 ` / ` 로, `-` 로 끝난 줄은 그냥 붙인다."""
    if v is None:
        return None
    out = ''
    for part in [p.strip() for p in str(v).split('\n') if p.strip()]:
        out += part if (not out or out.endswith('-')) else ' / ' + part
    return out or None


def status_of(text):
    if '단선' in text:
        return 'broken'
    if '사용불가' in text:
        return 'unusable'
    if '측정불가' in text:
        return 'ok'
    return None


def role_of(sheet, no, use):
    """회선수용내역 셀의 명시 표기 → 'main' | 'backup' | None. 둘 다면 실패."""
    if use is None:
        return None
    m, b = ROLE_MAIN.search(use), ROLE_BACKUP.search(use)
    if m and b:
        fail(f'{sheet} No.{no}: 주·예비 표기가 한 셀에 함께 있다 「{use}」')
    return 'main' if m else 'backup' if b else None


def plates_of(ws):
    """시트 → [{ peer, line, rows: [(no, row)] }] — 상대국명이 새로 적힌 코어 행이 판의 시작."""
    plates, cur, in_block = [], None, False
    for row in ws.iter_rows(values_only=True):
        row = list(row) + [None] * 16
        no = row[COL_NO]
        if no == 'No.':
            if [clean(x) for x in row[:5]] != HEADER:
                fail(f'{ws.title}: 열 머리가 다르다 {row[:5]}')
            in_block, cur = True, None
            continue
        if not (isinstance(no, (int, float)) and not isinstance(no, bool)):
            in_block, cur = False, None
            continue
        if not in_block:
            fail(f'{ws.title}: 열 머리 없는 코어 행 No.{no}')
        if row[COL_PEER] is not None:
            cur = {'peer': clean(row[COL_PEER]), 'line': line_name(row[COL_LINE]), 'rows': []}
            plates.append(cur)
        if cur is None:
            fail(f'{ws.title}: 상대국명 없이 시작한 블록 No.{no}')
        if cur['rows'] and int(no) != cur['rows'][-1][0] + 1:
            fail(f'{ws.title}: No. 가 이어지지 않는다 {cur["rows"][-1][0]} → {no}')
        cur['rows'].append((int(no), row))
    return plates


def cores_of(sheet, plate):
    start = plate['rows'][0][0]
    out = {}
    ranges = []
    for no, row in plate['rows']:
        core = no - start + 1
        kind, use, note, result = (clean(row[c]) for c in (COL_KIND, COL_USE, COL_NOTE, COL_RESULT))
        if note is not None and plate['line'] is not None and note == clean(plate['line'].replace(' / ', ' ')):
            note = None
        status, status_texts, purpose = 'ok', [], []
        for cell in (kind, use):
            if cell is None:
                continue
            st = status_of(cell)
            if st is None:
                purpose.append(cell)
            else:
                status_texts.append(cell)
                if st != 'ok' and status == 'ok':
                    status = st
        if note is not None:
            st = status_of(note)
            if st not in (None, 'ok') and status == 'ok':
                status = st
            m = RANGE.search(note)
            if m and st == 'unusable':
                ranges.append((int(m.group(1)), int(m.group(2)), note))
        if result == '단선':
            status = 'broken'
            status_texts.append(result)
        notes = status_texts + ([note] if note is not None else [])
        out[core] = {'no': core, 'status': status, 'role': role_of(sheet, no, use),
                     'purpose': ' · '.join(purpose) or None, 'note': ' · '.join(notes) or None}
    for lo, hi, text in ranges:
        if not (1 <= lo <= hi <= len(plate['rows'])):
            fail(f'{sheet}: 블록 주석 범위 {lo}~{hi} 가 판 코어 수 {len(plate["rows"])} 밖')
        for core in range(lo, hi + 1):
            c = out[core]
            c['status'] = 'unusable'
            if c['note'] is None or text not in c['note']:
                c['note'] = ' · '.join(x for x in (text, c['note']) if x)
    return [c for c in out.values() if c['status'] != 'ok' or c['purpose'] or c['note']]


def ts_plates(path):
    """`jikhalGangwon.ts` 의 PLATES — { key: { slot: (cores, peer) } }."""
    src = open(path, encoding='utf-8').read()
    body = src[src.index('export const PLATES'):]
    body = body[:body.index('\n};')]
    out, key = {}, None
    for line in body.split('\n'):
        m = re.match(r'\s+(\w+): \[', line)
        if m:
            key = m.group(1)
            out[key] = {}
        for s, c, p in re.findall(r"\{ slot: (\d+), cores: (\d+), peer: '([^']*)' \}", line):
            out[key][int(s)] = (int(c), p)
    return out


def ts_plate_lines(path):
    """`jikhalGangwonLines.ts` 의 PLATE_LINES — { key: { slot: line|None } }."""
    src = open(path, encoding='utf-8').read()
    body = src[src.index('export const PLATE_LINES'):]
    out, key, slot = {}, None, None
    for line in body.split('\n'):
        m = re.match(r'  (\w+): \{', line)
        if m:
            key = m.group(1)
            out[key] = {}
            continue
        m = re.match(r'    (\d+): \{', line)
        if m:
            slot = int(m.group(1))
            out[key][slot] = None
            continue
        m = re.search(r"line: '([^']*)'", line)
        if m and slot is not None:
            out[key][slot] = m.group(1)
    return out


def main(xlsx, out_path):
    seed_dir = out_path.rsplit('/', 1)[0]
    plates_ts = ts_plates(f'{seed_dir}/jikhalGangwon.ts')
    lines_ts = ts_plate_lines(f'{seed_dir}/jikhalGangwonLines.ts')
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    result = {}
    for sheet, key in SHEET_KEYS.items():
        if sheet not in wb.sheetnames:
            fail(f'시트 「{sheet}」 가 없다')
        plates = plates_of(wb[sheet])
        want = plates_ts.get(key)
        if want is None:
            fail(f'PLATES 에 국 {key} 가 없다')
        if len(plates) != len(want):
            fail(f'{sheet}: 판 {len(plates)}장 ≠ PLATES {len(want)}장')
        result[key] = {}
        for slot, plate in enumerate(plates, 1):
            cores, peer = want[slot]
            if (len(plate['rows']), plate['peer']) != (cores, peer):
                fail(f'{sheet} 판 {slot}: ({len(plate["rows"])}, {plate["peer"]}) ≠ PLATES ({cores}, {peer})')
            if slot in lines_ts.get(key, {}):
                if lines_ts[key][slot] != plate['line']:
                    fail(f'{sheet} 판 {slot}: 선로명 「{plate["line"]}」 ≠ PLATE_LINES 「{lines_ts[key][slot]}」')
            elif plate['line'] is not None:
                fail(f'{sheet} 판 {slot}: 선로명 「{plate["line"]}」 이 PLATE_LINES 에 없다')
            got = cores_of(sheet, plate)
            if got:
                result[key][slot] = got
    write(result, out_path)
    flat = [c for st in result.values() for cs in st.values() for c in cs]
    print(f'[extract_ofd_cores] 코어 {len(flat)}: 용도 {sum(1 for c in flat if c["purpose"])}, '
          f'단선 {sum(1 for c in flat if c["status"] == "broken")}, 사용불가 {sum(1 for c in flat if c["status"] == "unusable")}, '
          f'비고 {sum(1 for c in flat if c["note"])}, 주 {sum(1 for c in flat if c["role"] == "main")}, '
          f'예비 {sum(1 for c in flat if c["role"] == "backup")} → {out_path}')


def ts_str(v):
    return 'null' if v is None else "'" + v.replace('\\', '\\\\').replace("'", "\\'") + "'"


def write(result, out_path):
    out = [
        '/**',
        ' * 강원본부 직할 OFD 판의 **코어 용도·상태·비고** — `직할/직할_OFD선번장_260203.xlsx` 에서 뽑았다.',
        ' *',
        ' * 생성물이다 — **손으로 고치지 않는다**. 선번장을 다시 받으면 다시 뽑는다:',
        ' *   python3 -I scripts/extract-jikhal/extract_ofd_cores.py 직할/직할_OFD선번장_260203.xlsx backend/src/seed/jikhalGangwonCores.ts',
        ' *',
        ' * 키: 국 → 판 번호(`PLATES`·`PLATE_LINES` 와 같은 선번장 판 번호) → 판 안 코어(1부터). 규칙은 스크립트 머리말.',
        ' * 값이 모두 비고 상태가 정상인 코어는 싣지 않는다. role 은 회선수용내역의 명시 표기(㈜·(예)·주회선·예비회선 …)만.',
        ' */',
        "export interface PlateCore { no: number; status: 'ok' | 'broken' | 'unusable'; role: 'main' | 'backup' | null; purpose: string | null; note: string | null }",
        '',
        'export const PLATE_CORES: Record<string, Record<number, PlateCore[]>> = {',
    ]
    for key, plates in result.items():
        out.append(f'  {key}: {{')
        for slot, cores in plates.items():
            out.append(f'    {slot}: [')
            for c in cores:
                role = 'null' if c['role'] is None else f"'{c['role']}'"
                out.append(f"      {{ no: {c['no']}, status: '{c['status']}', role: {role}, purpose: {ts_str(c['purpose'])}, note: {ts_str(c['note'])} }},")
            out.append('    ],')
        out.append('  },')
    out.append('};')
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(out) + '\n')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit('usage: extract_ofd_cores.py <선번장.xlsx> <out.ts>')
    main(sys.argv[1], sys.argv[2])
